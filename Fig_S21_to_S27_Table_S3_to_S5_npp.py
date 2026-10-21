import os
import re
import sys
import glob
import time
import datetime
import numpy as np
import pandas as pd
import rasterio
from rasterio.features import rasterize
from scipy.stats import linregress
import geopandas as gpd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.colors as mcolors
import matplotlib.patches as patches
import matplotlib.patheffects as path_effects
import matplotlib.ticker as mticker
from matplotlib.colors import BoundaryNorm, LinearSegmentedColormap, ListedColormap
from matplotlib.lines import Line2D

try:
    import cartopy.crs as ccrs
    import cartopy.feature as cfeature
    HAVE_CARTOPY = True
except Exception:
    HAVE_CARTOPY = False

YEARS = [2013, 2014, 2015, 2016, 2017]
REF_YEAR = 2017
FIG_NUMBERS = {2013: 22, 2014: 23, 2015: 24, 2016: 25, 2017: 26}
NPP_PATH = "data/npp/npp_yearmean_rate_{year}.tif"
EMISSION_PATH = "data/emission/NO_{year}_yearly.tif"
CHL_WINDOWS = {y: f"surrogate_102_{y}/windows/windows_baseline.csv" for y in YEARS}
CHL_RESULTS = {y: f"surrogate_102_{y}/results_{y}.csv" for y in YEARS}
CHL_PATH = "data/chlor/chlor_{year}.tif"
NODATA_VALUES = [-9999.0, -999.0]
AOI_LIST = [("AOI 1", "NW Pacific off Japan", "all_AOIs_geojson/AOI1_NW_Pacific_off_Japan.geojson"),
            ("AOI 2", "Central South Atlantic", "all_AOIs_geojson/AOI2_Central_South_Atlantic.geojson"),
            ("AOI 3", "North Atlantic", "all_AOIs_geojson/AOI3_North_Atlantic.geojson"),
            ("AOI 4", "NE Pacific", "all_AOIs_geojson/AOI4_NE_Pacific.geojson"),
            ("AOI 5", "Eastern Pacific", "all_AOIs_geojson/AOI5_Eastern_Pacific.geojson")]
POOLED = "AOI 1 to 4 pooled"
OUT_DIR = os.path.join("results", "Fig_S21_to_S27_Table_S3_to_S5")
WIN_DIR = os.path.join(OUT_DIR, "windows")

R = 6_371_000
WINDOW_SIZE = 11
MIN_VALID = 70
R2_SCREEN = 0.6
SLOPE_ABS_MIN = 0.01

FONT = "Times New Roman"
FONT_SIZE = 13
FORMATS = ("png", "pdf")
R_LOW = np.sqrt(R2_SCREEN)
R_COLORS = ["#0000cc", "#8080ff", "#b3b3ff", "#ffb3b3", "#ff8080", "#cc0000"]
LAND_COLOR = "white"
LANE_SHP = "data/shipping_lane/mygeodata/Shipping_Lanes_v1.shp"
LANE_COLOR = "#7b3294"
LANE_TYPE_FIELD = "Type"
LANE_STYLE = {"core": {"linewidth": 0.8, "alpha": 0.9}, "secondary": {"linewidth": 0.5, "alpha": 0.6},
              "minor": {"linewidth": 0.3, "alpha": 0.4}, "default": {"linewidth": 0.3, "alpha": 0.4}}
LANE_CLASS_ALIASES = {"major": "core", "main": "core", "middle": "secondary", "medium": "secondary", "small": "minor"}
LANE_LEGEND = [("core", "Core Shipping Route"), ("secondary", "Secondary Shipping Route"), ("minor", "Minor Shipping Route")]
N_SHP = "data/n_limited_shp/Nmean60m_all_from_csv.shp"
N_FIELD = "meanN60"
N_CELL = 1.0
N_STRETCH = "histeq_gamma"
N_GAMMA = 1.0
N_DARK = "#a6a6a6"
N_LABEL = "Mean nitrate concentration\nin the upper 60 m (µmol kg⁻¹)"

CLS_JOINT, CLS_SILENT, CLS_DISAGREE, CLS_NPPONLY = 1, 2, 3, 4
CLS_COLOR = {CLS_JOINT: "#D55E00", CLS_SILENT: "#F0E442", CLS_DISAGREE: "#0072B2", CLS_NPPONLY: "#009E73"}
CLS_LABEL = {CLS_JOINT: "Positive in both Chl $a$ and NPP", CLS_SILENT: "Chl $a$ positive, NPP not retained",
             CLS_DISAGREE: "Chl $a$ positive, NPP negative", CLS_NPPONLY: "Positive in NPP only"}
AOI_EDGE_COLOR = "black"
AOI_LINESTYLE = (0, (4, 3))
AOI_LINEWIDTH = 1.0
AOI_LABEL_AT = {}
BASE_FONT = 14
FIGSIZE_S21 = (12.0, 6.8)

NODATA = 99999.0
PIXEL_RULE = "all-valid"
M_C, M_N, M_NO = 12.011, 14.007, 30.006
G_N_PER_G_NO = M_N / M_NO
CN_PER_SLOPE = M_NO / M_C
REDFIELD = 106.0 / 16.0
DELIVERY = 4.5 / 6.0
ANCHORS = {POOLED: (46.9, 1.0, 0.01, 0.005), "AOI 5": (1109.8, 5.0, 0.51, 0.01)}
AOI_COLORS = {"AOI 1": "#1f77b4", "AOI 2": "#ff7f0e", "AOI 3": "#2ca02c", "AOI 4": "#9467bd", "AOI 5": "#8c564b"}
HEADER_S5 = ["Region", "n (grid-cell years)", "Slope (g C m⁻² yr⁻¹ per g NO m⁻² yr⁻¹)", "95% CI half-width",
             "R²", "p", "Slope per g N (g C per g N)", "Molar C:N, emitted N", "Implied f, emitted N",
             "Molar C:N, deposited N", "Implied f, deposited N"]

plt.rcParams.update({"font.family": FONT, "font.size": FONT_SIZE, "axes.titlesize": FONT_SIZE,
                     "legend.fontsize": FONT_SIZE, "pdf.fonttype": 42, "ps.fonttype": 42,
                     "mathtext.fontset": "custom", "mathtext.rm": FONT, "mathtext.it": f"{FONT}:italic",
                     "mathtext.bf": f"{FONT}:bold"})
NOX = "NO$_{\\mathrm{x}}$"


def pixel_area_m2(lat_deg, xres_deg, yres_deg):
    lat = np.deg2rad(lat_deg); dlat = np.deg2rad(yres_deg); dlon = np.deg2rad(xres_deg)
    return R**2 * dlon * (np.sin(lat + dlat/2) - np.sin(lat - dlat/2))


def perform_moving_window_regression(fld, em, tf, window_size=WINDOW_SIZE):
    """Regress log1p field on log1p emission in every window; windows below SLOPE_ABS_MIN are kept and flagged floor_ok = False."""
    rows, cols = fld.shape
    half = window_size // 2
    with np.errstate(invalid="ignore", divide="ignore", over="ignore"):
        keep = np.isfinite(fld) & (fld > 0) & np.isfinite(em) & (em > 0)
        la, le = np.log1p(fld), np.log1p(em)
    results = []
    for i in range(rows):
        r0, r1 = max(i - half, 0), min(i + half + 1, rows)
        for j in range(cols):
            c0, c1 = max(j - half, 0), min(j + half + 1, cols)
            k = keep[r0:r1, c0:c1].ravel()
            n = int(k.sum())
            if n < MIN_VALID:
                continue
            a_vals = la[r0:r1, c0:c1].ravel()[k]
            b_vals = le[r0:r1, c0:c1].ravel()[k]
            if np.all(a_vals == a_vals[0]) or np.all(b_vals == b_vals[0]):
                continue
            slope, intercept, r_value, p_value, _ = linregress(b_vals, a_vals)
            lon, lat = tf * (j + 0.5, i + 0.5)
            results.append({"row": i, "col": j, "Longitude": lon, "Latitude": lat, "Intercept": intercept,
                            "Slope": slope, "R": r_value, "R2": r_value**2, "P_Value": p_value, "Count": n,
                            "floor_ok": bool(abs(slope) >= SLOPE_ABS_MIN)})
    return pd.DataFrame(results)


def read_masked(path):
    with rasterio.open(path) as src:
        arr = src.read(1).astype("float64"); nod = src.nodata; tf = src.transform; shape = src.shape; xres, yres = src.res
    if nod is not None and np.isfinite(nod):
        arr[arr == nod] = np.nan
    return arr, tf, shape, xres, yres


def cell_key(lon, lat, xres, yres):
    return (np.round((90 - lat) / yres - 0.5).astype(int), np.round((lon + 180) / xres - 0.5).astype(int))


def load_chl_screened(year, xres, yres):
    """Chlorophyll a retained windows of one year with row and col, checked against the surrogate baseline."""
    p = CHL_WINDOWS[year]
    if not os.path.exists(p):
        sys.exit(f"[ABORT] chlorophyll window file not found: {p}")
    df = pd.read_csv(p); df = df[df["R2"] >= R2_SCREEN].reset_index(drop=True)
    ref = pd.read_csv(CHL_RESULTS[year]); ref = int(ref[ref["scenario"] == "baseline"].iloc[-1]["n_screened"])
    if len(df) != ref:
        sys.exit(f"[ABORT] {p}: {len(df):,} retained windows, the surrogate baseline gives {ref:,}")
    df["row"], df["col"] = cell_key(df["Longitude"].to_numpy(), df["Latitude"].to_numpy(), xres, yres)
    return df


def aoi_masks(ref_path):
    """Boolean mask per AOI on the reference grid (all_touched), the pooled mask and the AOI index grid (first AOI wins, -1 outside)."""
    masks = {}
    with rasterio.open(ref_path) as src:
        for short, _, gj in AOI_LIST:
            if not os.path.exists(gj):
                sys.exit(f"[ABORT] missing input {gj}")
            gdf = gpd.read_file(gj)
            if gdf.crs is not None and src.crs is not None and gdf.crs != src.crs:
                gdf = gdf.to_crs(src.crs)
            shapes = [g for g in gdf.geometry if g is not None and g.is_valid and g.geom_type in ("Polygon", "MultiPolygon")]
            masks[short] = rasterize([(g, 1) for g in shapes], out_shape=(src.height, src.width), transform=src.transform,
                                     fill=0, dtype="uint8", all_touched=True).astype(bool)
            print(f"[AOI] {short}: {int(masks[short].sum())} grid cells")
    masks[POOLED] = np.logical_or.reduce([masks[s] for s, _, _ in AOI_LIST[:4]])
    grid = np.full(masks["AOI 1"].shape, -1, dtype=int)
    for k, (short, _, _) in reversed(list(enumerate(AOI_LIST))):
        grid[masks[short]] = k
    return masks, grid


def have_natural_earth():
    """True when cartopy can supply the land and coastline features."""
    if not HAVE_CARTOPY:
        print("[note] cartopy not available; drawing plain longitude-latitude frames"); return False
    try:
        list(cfeature.LAND.geometries()); list(cfeature.COASTLINE.geometries()); return True
    except Exception as e:
        print(f"[note] coastline data not available ({type(e).__name__}); land and coastlines skipped"); return False


def lane_class(value):
    v = str(value).strip().lower(); return LANE_CLASS_ALIASES.get(v, v)


def load_lane_shapes():
    """The shipping-route shapefile in WGS84, or None when it is absent."""
    if not os.path.exists(LANE_SHP):
        print(f"[warn] routes shapefile not found ({LANE_SHP}); maps drawn without routes"); return None
    lanes = gpd.read_file(LANE_SHP)
    if lanes.crs is not None and lanes.crs.to_epsg() != 4326:
        lanes = lanes.to_crs(4326)
    print(f"[lanes] {len(lanes)} features; classes {lanes[LANE_TYPE_FIELD].value_counts().to_dict() if LANE_TYPE_FIELD in lanes.columns else 'n/a'}")
    return lanes


def draw_lanes(ax, lanes, zorder):
    if LANE_TYPE_FIELD in lanes.columns:
        for cls, g in lanes.groupby(LANE_TYPE_FIELD):
            st = LANE_STYLE.get(lane_class(cls), LANE_STYLE["default"])
            ax.add_geometries(g.geometry, crs=ccrs.PlateCarree(), facecolor="none", edgecolor=LANE_COLOR, zorder=zorder, **st)
    else:
        ax.add_geometries(lanes.geometry, crs=ccrs.PlateCarree(), facecolor="none", edgecolor=LANE_COLOR, zorder=zorder, **LANE_STYLE["default"])


def r_class(r):
    """Index of the six r classes: three per sign, split at 0.8 and 0.9."""
    if r < 0:
        return 0 if r <= -0.9 else 1 if r <= -0.8 else 2
    return 5 if r >= 0.9 else 4 if r >= 0.8 else 3


def load_n_grid():
    """Rasterise the WOA23 nitrate point shapefile on a 1-degree grid; returns (grid, lon_edges, lat_edges, max)."""
    if not os.path.exists(N_SHP):
        sys.exit(f"[ABORT] nitrate shapefile not found: {N_SHP}")
    pts = gpd.read_file(N_SHP)
    if pts.crs is not None and pts.crs.to_epsg() != 4326:
        pts = pts.to_crs(4326)
    if N_FIELD not in pts.columns:
        sys.exit(f"[ABORT] {N_SHP}: field '{N_FIELD}' not found in {list(pts.columns)}")
    pts = pts[pts.geometry.notna() & pts[N_FIELD].notna()]
    lon = pts.geometry.x.to_numpy(float); lat = pts.geometry.y.to_numpy(float); val = pts[N_FIELD].to_numpy(float)
    ncols, nrows = int(round(360 / N_CELL)), int(round(180 / N_CELL))
    jj = np.clip(np.floor((lon + 180) / N_CELL).astype(int), 0, ncols - 1)
    ii = np.clip(np.floor((90 - lat) / N_CELL).astype(int), 0, nrows - 1)
    acc = np.zeros((nrows, ncols)); cnt = np.zeros((nrows, ncols), dtype=int)
    np.add.at(acc, (ii, jj), val); np.add.at(cnt, (ii, jj), 1)
    grid = np.full((nrows, ncols), np.nan); ok = cnt > 0; grid[ok] = acc[ok] / cnt[ok]
    if (cnt > 1).any():
        print(f"[N] note: {int((cnt > 1).sum())} cells hold more than one point; their mean is used")
    lon_e = -180 + np.arange(ncols + 1) * N_CELL; lat_e = 90 - np.arange(nrows + 1) * N_CELL
    print(f"[N] {N_SHP}: {len(pts):,} points, field '{N_FIELD}', {N_CELL:g}-degree grid, {int(ok.sum()):,} cells, max {np.nanmax(grid):.3f}")
    return grid, lon_e, lat_e, float(np.nanmax(grid))


def n_stretched(grid):
    """Map the nitrate field to [0, 1] for the grey ramp (histogram equalisation, gamma or linear)."""
    out = np.full(grid.shape, np.nan); ok = np.isfinite(grid)
    if N_STRETCH == "linear":
        out[ok] = grid[ok] / np.nanmax(grid)
    elif N_STRETCH == "gamma":
        out[ok] = (grid[ok] / np.nanmax(grid)) ** N_GAMMA
    else:
        v = np.sort(grid[ok]); cdf = np.arange(1, len(v) + 1) / len(v)
        out[ok] = np.interp(grid[ok], v, cdf) ** N_GAMMA
    return out


def draw_legend_strip(fig, n_cmap, n_max, lanes_present, fp):
    """Legend strip below the map: r blocks with rotated labels, the nitrate ramp and the three route classes."""
    lax = fig.add_axes([0.03, 0.055, 0.94, 0.155]); lax.set_xlim(0, 1); lax.set_ylim(0, 1); lax.axis("off")
    top = 0.92
    low = f"{R_LOW:.2f}"
    blocks = [(R_COLORS[3], f"{low} - 0.8"), (R_COLORS[4], "0.8 - 0.9"), (R_COLORS[5], "> 0.9"), None,
              (R_COLORS[2], f"-0.8 - -{low}"), (R_COLORS[1], "-0.9 - -0.8"), (R_COLORS[0], "< -0.9")]
    x, w, h = 0.04, 0.045, 0.17
    for b in blocks:
        if b is None:
            x += 0.035; continue
        color, label = b
        lax.add_patch(patches.Rectangle((x, top - h), w, h, facecolor=color, edgecolor="none"))
        lax.text(x + w / 2, top - h - 0.06, label, ha="right", va="top", rotation=45, rotation_mode="anchor", **fp)
        x += w
    xn = 0.50
    lax.text(xn, top, N_LABEL, ha="left", va="top", linespacing=1.3, **fp)
    bx, by, bw, bh = xn + 0.02, 0.06, 0.035, 0.42
    lax.imshow(np.linspace(0, 1, 64)[:, None], extent=(bx, bx + bw, by, by + bh), origin="lower", cmap=n_cmap,
               aspect="auto", zorder=2, vmin=0, vmax=1)
    lax.add_patch(patches.Rectangle((bx, by), bw, bh, facecolor="none", edgecolor="#999999", linewidth=0.4, zorder=3))
    lax.text(bx + bw + 0.008, by + bh, f"{n_max:.1f}", ha="left", va="center", **fp)
    lax.text(bx + bw + 0.008, by, "0", ha="left", va="center", **fp)
    if lanes_present:
        for i, (cls, label) in enumerate(LANE_LEGEND):
            yy = top - 0.04 - i * 0.24
            lax.plot([0.72, 0.765], [yy, yy], color=LANE_COLOR, **LANE_STYLE[cls])
            lax.text(0.775, yy, label, ha="left", va="center", **fp)


def draw_r_map(scr, title, rows, cols, n_layers, lanes, out_base, have_ne):
    """One window map: r classes of the screened windows (columns row, col, R) over the nitrate basemap and the routes."""
    n_grid, n_lon_e, n_lat_e, n_max = n_layers
    n_cmap = LinearSegmentedColormap.from_list("n60", ["#ffffff", N_DARK])
    field = np.full((rows, cols), np.nan)
    for r_, c_, rv in zip(scr["row"], scr["col"], scr["R"]):
        field[int(r_), int(c_)] = r_class(float(rv))
    lon_e = np.linspace(-180, 180, cols + 1); lat_e = np.linspace(90, -90, rows + 1)
    fp = {"family": FONT, "size": FONT_SIZE}
    fig = plt.figure(figsize=(12, 7.8))
    if HAVE_CARTOPY:
        ax = fig.add_axes([0.02, 0.235, 0.96, 0.755], projection=ccrs.Robinson()); ax.set_global(); kw = {"transform": ccrs.PlateCarree()}
    else:
        ax = fig.add_axes([0.04, 0.235, 0.92, 0.755]); ax.set_xlim(-180, 180); ax.set_ylim(-90, 90); ax.set_aspect("equal"); kw = {}
    ax.pcolormesh(n_lon_e, n_lat_e, np.ma.masked_invalid(n_stretched(n_grid)), cmap=n_cmap, vmin=0, vmax=1, shading="flat", zorder=1, **kw)
    if HAVE_CARTOPY and have_ne:
        ax.add_feature(cfeature.LAND, facecolor=LAND_COLOR, zorder=2)
        ax.add_feature(cfeature.COASTLINE, linewidth=0.4, zorder=5)
    if HAVE_CARTOPY and lanes is not None:
        draw_lanes(ax, lanes, zorder=3)
    cmap = ListedColormap(R_COLORS); norm = BoundaryNorm(np.arange(-0.5, 6.5, 1), cmap.N)
    ax.pcolormesh(lon_e, lat_e, np.ma.masked_invalid(field), cmap=cmap, norm=norm, shading="flat", zorder=4, **kw)
    if HAVE_CARTOPY:
        gl = ax.gridlines(draw_labels=True, xlocs=range(-180, 181, 60), ylocs=[-60, -30, 0, 30, 60],
                          linewidth=0.5, color="gray", alpha=0.5, linestyle="--", zorder=6)
        gl.top_labels = False; gl.xlabel_style = dict(fp); gl.ylabel_style = dict(fp)
    ax.set_title(title, fontsize=FONT_SIZE, fontfamily=FONT)
    draw_legend_strip(fig, n_cmap, n_max, lanes is not None, fp)
    for ext in FORMATS:
        fig.savefig(f"{out_base}.{ext}", dpi=300, bbox_inches="tight")
    plt.close(fig)


def add_base_features(ax, coast_lw=0.4):
    for feat, kw in [(cfeature.LAND, {"facecolor": "lightgray"}), (cfeature.OCEAN, {"facecolor": "white"}), (cfeature.COASTLINE, {"linewidth": coast_lw})]:
        try:
            next(iter(feat.geometries()), None)
            ax.add_feature(feat, **kw)
        except Exception as e:
            print(f"[warn] base feature '{getattr(feat, 'name', feat)}' skipped: {e}")


def add_gridlines(ax, label_size, lon_step=60):
    kw = dict(linewidth=0.4, color="gray", alpha=0.55, linestyle=":",
              xlocs=mticker.FixedLocator(range(-180, 181, lon_step)), ylocs=mticker.FixedLocator(range(-90, 91, 30)))
    try:
        gl = ax.gridlines(draw_labels=True, **kw)
        gl.top_labels = False; gl.right_labels = False
        gl.xlabel_style = {"size": label_size}; gl.ylabel_style = {"size": label_size}
    except Exception as e:
        print(f"[warn] graticule labels unavailable ({e}); lines only"); ax.gridlines(draw_labels=False, **kw)


def draw_aois(ax, aoi_gdfs):
    for name, gdf in aoi_gdfs.items():
        ax.add_geometries(gdf.geometry, crs=ccrs.PlateCarree(), facecolor="none", edgecolor=AOI_EDGE_COLOR,
                          linestyle=AOI_LINESTYLE, linewidth=AOI_LINEWIDTH, zorder=6)
        if name in AOI_LABEL_AT:
            lo, la = AOI_LABEL_AT[name]
        else:
            minx, miny, maxx, maxy = gdf.total_bounds
            lo, la = minx, maxy + 1.5
        ax.text(lo, la, name, fontsize=BASE_FONT, color="black", ha="left", va="bottom", zorder=7, transform=ccrs.PlateCarree(),
                path_effects=[path_effects.withStroke(linewidth=2.0, foreground="white")])


def draw_overlap_map(cells_by_class, out_base, shape, lanes, aoi_gdfs):
    """Fig. S21: the four overlap classes of 2017 with the routes and the AOIs."""
    rows, cols = shape
    grid = np.full(shape, np.nan)
    for cls_value, cells in cells_by_class.items():
        for r_, c_ in cells:
            grid[r_, c_] = cls_value
    xs = np.linspace(-180, 180, cols + 1); ys = np.linspace(90, -90, rows + 1)
    cmap = mcolors.ListedColormap([CLS_COLOR[c] for c in (CLS_JOINT, CLS_SILENT, CLS_DISAGREE, CLS_NPPONLY)])
    norm = mcolors.BoundaryNorm([0.5, 1.5, 2.5, 3.5, 4.5], cmap.N)
    fig = plt.figure(figsize=FIGSIZE_S21)
    ax = fig.add_subplot(1, 1, 1, projection=ccrs.Robinson()); ax.set_global()
    add_base_features(ax)
    add_gridlines(ax, BASE_FONT)
    ax.pcolormesh(xs, ys, grid, transform=ccrs.PlateCarree(), cmap=cmap, norm=norm, shading="flat")
    if lanes is not None:
        draw_lanes(ax, lanes, zorder=5)
    draw_aois(ax, aoi_gdfs)
    handles = [Line2D([0], [0], marker="s", color=CLS_COLOR[c], linestyle="None", markersize=11, label=CLS_LABEL[c])
               for c in (CLS_JOINT, CLS_SILENT, CLS_DISAGREE, CLS_NPPONLY)]
    handles += [Line2D([0], [0], color=LANE_COLOR, lw=1.8, alpha=0.9, label="core shipping routes"),
                Line2D([0], [0], color=LANE_COLOR, lw=1.1, alpha=0.6, label="secondary shipping routes"),
                Line2D([0], [0], color=LANE_COLOR, lw=0.6, alpha=0.4, label="minor shipping routes"),
                Line2D([0], [0], color=AOI_EDGE_COLOR, ls="--", lw=1.2, label="AOI boundaries")]
    fig.legend(handles=handles, loc="lower center", ncol=4, frameon=False, fontsize=BASE_FONT)
    fig.subplots_adjust(left=0.03, right=0.97, top=0.98, bottom=0.16)
    for ext in FORMATS:
        fig.savefig(f"{out_base}.{ext}", dpi=300, bbox_inches="tight")
    plt.close(fig)
    print(f"[Saved] {out_base}.{FORMATS[0]}")


def n_pct(n, den):
    return f"{n:,} ({n / den * 100:.1f}%)" if den else f"{n:,} (0.0%)"


def table_s3(sets, aoi_grid, npp_scr):
    """Table S3 (A): the overlap classes by region (global, the AOIs combined, each AOI, outside the AOIs)."""
    cpos, joint, disagree, silent, npp_only = (sets[k] for k in ("cpos", "joint", "disagree", "silent", "npp_only"))
    npp_aoi = aoi_grid[npp_scr["row"].to_numpy(), npp_scr["col"].to_numpy()]; npp_pos = npp_scr["Slope"].to_numpy() > 0
    scopes = [("Global ocean", lambda k: True, np.ones(len(npp_scr), bool)),
              ("All 5 AOIs combined", lambda k: aoi_grid[k] >= 0, npp_aoi >= 0)]
    for i, (short, long_, _) in enumerate(AOI_LIST):
        scopes.append((f"{short} {long_}", lambda k, i=i: aoi_grid[k] == i, npp_aoi == i))
    scopes.append(("Outside the AOIs", lambda k: aoi_grid[k] < 0, npp_aoi < 0))
    rows = []
    for label, keep, m in scopes:
        cp = [c for c in cpos if keep(c)]
        rows.append({"Region": label, "Chl a-positive cells": f"{len(cp):,}",
                     "Positive in both": n_pct(sum(keep(c) for c in joint), len(cp)),
                     "Chl a positive, NPP negative": n_pct(sum(keep(c) for c in disagree), len(cp)),
                     "Chl a positive, NPP not retained": n_pct(sum(keep(c) for c in silent), len(cp)),
                     "Positive in NPP only": f"{sum(keep(c) for c in npp_only):,}",
                     "NPP windows retained (positive / total)": f"{int((m & npp_pos).sum()):,} / {int(m.sum()):,}"})
    return pd.DataFrame(rows)


def load_field(path):
    """2D float array with NaN for the nodata tag and the NODATA_VALUES."""
    with rasterio.open(path) as src:
        a = src.read(1).astype(float)
        if src.nodata is not None:
            a[a == src.nodata] = np.nan
    for v in NODATA_VALUES:
        a[a == v] = np.nan
    return a


def window_stats(chl, npp, emi):
    """Per-window statistics of chlorophyll a and NPP on the cells valid in all three fields (interior windows only)."""
    ny, nx = chl.shape
    h = WINDOW_SIZE // 2
    lchl, lnpp, lemi = np.log1p(chl), np.log1p(npp), np.log1p(emi)
    valid_all = np.isfinite(chl) & np.isfinite(npp) & np.isfinite(emi) & (chl >= 0) & (npp >= 0) & (emi >= 0)
    rows = []
    for i in range(h, ny - h):
        for j in range(h, nx - h):
            sl = (slice(i - h, i + h + 1), slice(j - h, j + h + 1))
            m = valid_all[sl]
            n = int(m.sum())
            if n < MIN_VALID:
                continue
            c, p = chl[sl][m], npp[sl][m]
            lc, lp, le = lchl[sl][m], lnpp[sl][m], lemi[sl][m]
            if np.var(le) == 0:
                continue
            def fit_y(y):
                xm, ym = le.mean(), y.mean()
                slope = ((le - xm) * (y - ym)).sum() / ((le - xm) ** 2).sum()
                sst = ((y - ym) ** 2).sum()
                r2 = 1.0 - ((y - (ym + slope * (le - xm))) ** 2).sum() / sst if sst > 0 else np.nan
                return slope, r2
            s_c, r_c = fit_y(lc); s_p, r_p = fit_y(lp)
            rows.append(dict(row=i, col=j, n_valid=n, mean_chl=c.mean(), mean_npp=p.mean(),
                             cv_chl=c.std(ddof=1) / c.mean() if c.mean() > 0 else np.nan,
                             cv_npp=p.std(ddof=1) / p.mean() if p.mean() > 0 else np.nan,
                             sd_logchl=lc.std(ddof=1), sd_lognpp=lp.std(ddof=1), sd_logemi=le.std(ddof=1),
                             slope_chl=s_c, r2_chl=r_c, slope_npp=s_p, r2_npp=r_p))
    return pd.DataFrame(rows)


def table_s3b():
    """Table S3 (B): chlorophyll a against NPP within the chlorophyll a-positive windows of 2017 (medians with IQR, shares)."""
    chl = load_field(CHL_PATH.format(year=REF_YEAR)); npp = load_field(NPP_PATH.format(year=REF_YEAR)); emi = load_field(EMISSION_PATH.format(year=REF_YEAR))
    if not (chl.shape == npp.shape == emi.shape):
        sys.exit(f"[ABORT] grids differ: chl {chl.shape}, npp {npp.shape}, emi {emi.shape}")
    df = window_stats(chl, npp, emi)
    df["chl_retained_pos"] = (df.slope_chl >= SLOPE_ABS_MIN) & (df.r2_chl >= R2_SCREEN)
    df["npp_retained_pos"] = (df.slope_npp >= SLOPE_ABS_MIN) & (df.r2_npp >= R2_SCREEN)
    df.to_csv(os.path.join(WIN_DIR, f"window_stats_{REF_YEAR}.csv"), index=False)
    sel = df[df.chl_retained_pos]
    q = lambda col: sel[col].quantile([0.25, 0.5, 0.75])
    miqr = lambda col: f"{q(col)[0.5]:.3f} ({q(col)[0.25]:.3f} to {q(col)[0.75]:.3f})"
    rows = [["Median coefficient of variation within a window (IQR)", miqr("cv_chl"), miqr("cv_npp")],
            ["Median R\u00b2 of the window regression on emission (IQR)", miqr("r2_chl"), miqr("r2_npp")],
            ["Windows with a positive slope", "100%", f"{100 * (sel.slope_npp > 0).mean():.1f}%"],
            [f"Windows passing the screen (R\u00b2 \u2265 {R2_SCREEN}, slope \u2265 {SLOPE_ABS_MIN})", "100%", f"{100 * sel.npp_retained_pos.mean():.1f}%"],
            ["Windows", f"{len(sel):,}", f"{len(sel):,}"]]
    tab = pd.DataFrame(rows, columns=["", "Chl a", "NPP"])
    p = os.path.join(OUT_DIR, "Table_S3B_within_window_comparison.csv"); tab.to_csv(p, index=False, encoding="utf-8-sig")
    print(f"[Saved] {p}"); print(tab.to_string(index=False))
    lines = [f"Windows entering the screen (>= {MIN_VALID} valid cells in chl, NPP and emission, interior windows): {len(df):,}",
             f"Chlorophyll a retained with positive slope: {len(sel):,}",
             f"median SD of log1p chlorophyll a {q('sd_logchl')[0.5]:.3f}, of log1p NPP {q('sd_lognpp')[0.5]:.3f}",
             f"median slope chlorophyll a {q('slope_chl')[0.5]:.3f}, NPP {q('slope_npp')[0.5]:.3f} (transformed units)",
             f"median ratio CV_NPP / CV_chl {(sel.cv_npp / sel.cv_chl).median():.3f}; share with NPP R2 >= {R2_SCREEN}: {100 * (sel.r2_npp >= R2_SCREEN).mean():.1f}%",
             f"all windows entering: median CV_chl {df.cv_chl.median():.3f}, CV_npp {df.cv_npp.median():.3f}, R2_chl {df.r2_chl.median():.3f}, R2_npp {df.r2_npp.median():.3f}; "
             f"NPP retained with positive slope anywhere {int(df.npp_retained_pos.sum()):,}"]
    return lines


def write_table_s4(rows, out_csv, out_docx):
    """Paste-ready Table S4 as CSV and, when python-docx is installed, as a Word table."""
    n = lambda x: f"{int(round(x)):,}"
    f = lambda x, d: f"{x:.{d}f}"
    tab = pd.DataFrame([{"Year": r["year"], "Fitted": n(r["fitted"]), "Entering the screen": n(r["entering"]),
                         "Retained": n(r["retained"]), "Positive": n(r["positive"]), "Negative": n(r["negative"]),
                         "Screened-window rate": f(r["screened_window_rate"], 4),
                         "Positive-correlation area (10⁶ km²)": f(r["pos_area_km2"] / 1e6, 2),
                         "Chlorophyll a-positive cells": n(r["chl_positive_cells"]),
                         "Positive in both": f"{n(r['positive_in_both'])} ({100 * r['share_positive_in_both']:.1f}%)",
                         "Chl a positive, NPP negative": n(r["chl_pos_npp_neg"]),
                         "Chl a positive, NPP not retained": n(r["chl_pos_npp_none"]),
                         "Positive in NPP only": n(r["npp_only_positive"])} for r in rows])
    tab.to_csv(out_csv, index=False, encoding="utf-8-sig"); print(f"[Saved] {out_csv}")
    try:
        from docx import Document
        from docx.shared import Pt
        doc = Document(); st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(9)
        doc.add_paragraph("Table S4. NPP-NOx retained windows and their overlap with the chlorophyll a retained windows, 2013 to 2017 (caption of record in the SI file).")
        t = doc.add_table(rows=1, cols=len(tab.columns)); t.style = "Table Grid"
        for j, c in enumerate(tab.columns):
            t.rows[0].cells[j].text = c
        for _, r in tab.iterrows():
            cells = t.add_row().cells
            for j, c in enumerate(tab.columns):
                cells[j].text = str(r[c])
        doc.save(out_docx); print(f"[Saved] {out_docx}")
    except ImportError:
        print("[note] python-docx not installed; the Word version was skipped")
    return tab


def read_band(path):
    """Band 1 as float64 with the fill value set to NaN, plus the grid (shape, affine, crs)."""
    with rasterio.open(path) as src:
        a = src.read(1).astype("float64")
        grid = (a.shape, tuple(src.transform)[:6], src.crs)
    a[a == NODATA] = np.nan
    return a, grid


def same_grid(g1, g2):
    return g1[0] == g2[0] and all(abs(a - b) < 1e-9 for a, b in zip(g1[1], g2[1]))


def pixels(npp, emi, mask, rule):
    """Emission (x) and NPP (y) of all grid-cell years in a region under one pixel rule."""
    xs, ys = [], []
    for y in YEARS:
        ok = mask & np.isfinite(npp[y]) & (npp[y] > 0) & np.isfinite(emi[y]) & (emi[y] >= 0)
        if rule == "emis>0":
            ok &= emi[y] > 0
        xs.append(emi[y][ok]); ys.append(npp[y][ok])
    return np.concatenate(xs), np.concatenate(ys)


def fit(x, y):
    """Untransformed OLS of NPP on emission: n, slope, 95% CI half-width (1.96 s.e.), R2, nominal p."""
    r = linregress(x, y)
    return dict(n=len(x), slope=r.slope, ci=1.96 * r.stderr, r2=r.rvalue ** 2, p=r.pvalue)


def stoichiometry(slope):
    """(slope per g N, C:N emitted, f emitted, C:N deposited, f deposited); NaN for a negative slope."""
    if not (np.isfinite(slope) and slope > 0):
        return (np.nan,) * 5
    cn = slope * CN_PER_SLOPE
    return slope / G_N_PER_G_NO, cn, REDFIELD / cn, cn / DELIVERY, REDFIELD / cn * DELIVERY


def check_anchors(fits):
    """Stop unless PIXEL_RULE reproduces the reference fits in ANCHORS; report the other rule as a hint."""
    passed = {}
    for rule in ("all-valid", "emis>0"):
        ok = True
        for region, (s0, s_tol, r0, r_tol) in ANCHORS.items():
            f = fits[rule][region]
            hit = abs(f["slope"] - s0) <= s_tol and abs(f["r2"] - r0) <= r_tol
            print(f"[Check] {rule:9s} {region:18s} slope {f['slope']:8.1f} (reference {s0}), R2 {f['r2']:.3f} (reference {r0}) -> {'ok' if hit else 'MISS'}")
            ok &= hit
        passed[rule] = ok
    if not passed[PIXEL_RULE]:
        other = "emis>0" if PIXEL_RULE == "all-valid" else "all-valid"
        raise SystemExit(f"[ERROR] PIXEL_RULE={PIXEL_RULE!r} does not reproduce the reference fits; "
                         + (f"{other!r} does, set PIXEL_RULE accordingly" if passed[other] else "neither rule does, the inputs differ from those of the reference fits"))
    print(f"[Check] PIXEL_RULE={PIXEL_RULE!r} reproduces the reference fits")


def table_s5(masks):
    """Table S5: pooled fits of NPP on emission by AOI, 2013 to 2017, converted to the implied C:N and f-ratio."""
    f0 = lambda x: "" if not np.isfinite(x) else f"{int(round(x)):,}"
    f1 = lambda x: "" if not np.isfinite(x) else f"{x:.1f}".replace("-", "−")
    f3 = lambda x: "" if not np.isfinite(x) else f"{x:.3f}"
    fr = lambda x: "" if not np.isfinite(x) else f"{x:.0f}"
    f2r = lambda x: "" if not np.isfinite(x) else ("<0.01" if x < 0.005 else f"{x:.2f}")
    fp = lambda x: "" if not np.isfinite(x) else ("<0.001" if x < 0.001 else f"{x:.3f}")
    npp, emi, ref = {}, {}, None
    for y in YEARS:
        npp[y], g1 = read_band(NPP_PATH.format(year=y))
        emi[y], g2 = read_band(EMISSION_PATH.format(year=y))
        ref = ref or g1
        if not (same_grid(g1, g2) and same_grid(ref, g1)):
            sys.exit(f"[ABORT] grid mismatch in {y}; the pooling assumes one shared grid")
    order = [s for s, _, _ in AOI_LIST] + [POOLED]
    fits = {rule: {name: fit(*pixels(npp, emi, masks[name], rule)) for name in order} for rule in ("all-valid", "emis>0")}
    check_anchors(fits)
    rows = []
    for name in order:
        f = fits[PIXEL_RULE][name]
        per_n, cn_e, f_e, cn_d, f_d = stoichiometry(f["slope"])
        rows.append([name, f0(f["n"]), f1(f["slope"]), f1(f["ci"]), f2r(f["r2"]), fp(f["p"]), f1(per_n), fr(cn_e), f3(f_e), fr(cn_d), f3(f_d)])
    path = os.path.join(OUT_DIR, "Table_S5_npp_stoichiometry.csv")
    table = pd.DataFrame(rows, columns=HEADER_S5)
    table.to_csv(path, index=False, encoding="utf-8-sig")
    with pd.option_context("display.width", 220, "display.max_columns", 20):
        print(table.to_string(index=False))
    print(f"[Constants] g N per g NO {G_N_PER_G_NO:.4f}; molar C:N per unit slope {CN_PER_SLOPE:.4f}; Redfield {REDFIELD:.3f}; delivery {DELIVERY:.2f}")
    print(f"[Saved] {path}")
    return npp, emi


def fig_s27(npp, emi, masks):
    """Fig. S27: NPP against NOx emission, every grid-cell year of the five AOIs, with the OLS line fitted on AOI 1 to 4 pooled."""
    import statsmodels.api as sm
    pts = {short: pixels(npp, emi, masks[short], "emis>0") for short, _, _ in AOI_LIST}
    x_fit = np.concatenate([pts[s][0] for s, _, _ in AOI_LIST[:4]]); y_fit = np.concatenate([pts[s][1] for s, _, _ in AOI_LIST[:4]])
    m = sm.OLS(y_fit, sm.add_constant(x_fit)).fit(); ci = m.conf_int(0.05)
    print(f"[Fig. S27] OLS on AOI 1 to 4 pooled: n={len(x_fit):,}, slope={m.params[1]:.8g} (+/-{(ci[1, 1] - ci[1, 0]) / 2:.8g}), intercept={m.params[0]:.8g}, R2={m.rsquared:.4f}")
    x5, y5 = pts["AOI 5"]; m5 = sm.OLS(y5, sm.add_constant(x5)).fit(); ci5 = m5.conf_int(0.05)
    print(f"[Fig. S27] AOI 5: n={len(x5):,}, slope={m5.params[1]:.8g} (+/-{(ci5[1, 1] - ci5[1, 0]) / 2:.8g}), R2={m5.rsquared:.4f}")
    plt.rcParams.update({"axes.linewidth": 1.2})
    fig, ax = plt.subplots(figsize=(8, 6))
    for short, long_, _ in AOI_LIST:
        ax.scatter(pts[short][0], pts[short][1], s=12, alpha=0.6, color=AOI_COLORS[short], label=f"{short}, {long_}", linewidths=0)
    x_all = np.concatenate([pts[s][0] for s, _, _ in AOI_LIST]); xl = np.linspace(x_all.min(), x_all.max(), 200)
    ax.plot(xl, m.params[0] + m.params[1] * xl, color="red", lw=2.5, label="Linear fit on AOI 1 to 4")
    ax.set_xlabel("NO$_{\\mathrm{x}}$ Emissions (g m$^{-2}$ yr$^{-1}$)"); ax.set_ylabel("NPP (g C m$^{-2}$ yr$^{-1}$)")
    ax.grid(True, alpha=0.3); ax.legend(frameon=True)
    fig.tight_layout()
    out = os.path.join(OUT_DIR, "Fig_S27_npp_nox_aoi_scatter")
    for ext in FORMATS:
        fig.savefig(f"{out}.{ext}", dpi=300)
    plt.close(fig); print(f"[Saved] {out}.{FORMATS[0]}")


def main():
    os.makedirs(WIN_DIR, exist_ok=True)
    for y in YEARS:
        for p in (NPP_PATH.format(year=y), EMISSION_PATH.format(year=y), CHL_RESULTS[y]):
            if not os.path.exists(p):
                sys.exit(f"[ABORT] missing input {p}")
    n_layers = load_n_grid()
    lanes = load_lane_shapes()
    have_ne = have_natural_earth()
    masks, aoi_grid = aoi_masks(NPP_PATH.format(year=REF_YEAR))

    rows, checks = [], []
    for year in YEARS:
        npp, tf, shape, xres, yres = read_masked(NPP_PATH.format(year=year)); em, tfe, she, _, _ = read_masked(EMISSION_PATH.format(year=year))
        if shape != she or tf != tfe:
            sys.exit(f"[ABORT] {year}: NPP and emission rasters are not on one grid\n  {tf}\n  {tfe}")
        t0 = time.time()
        df = perform_moving_window_regression(npp, em, tf)
        df.to_csv(os.path.join(WIN_DIR, f"windows_npp_{year}_all.csv"), index=False, encoding="utf-8-sig")
        ent = df[df["floor_ok"]]; scr = ent[ent["R2"] >= R2_SCREEN]
        pos = scr[scr["Slope"] >= 0]; neg = scr[scr["Slope"] < 0]
        area = lambda d: float(pixel_area_m2(d["Latitude"].to_numpy(), xres, yres).sum() / 1e6) if len(d) else 0.0

        chl = load_chl_screened(year, xres, yres)
        chl_pos = set(zip(chl.loc[chl["Slope"] >= 0, "row"], chl.loc[chl["Slope"] >= 0, "col"]))
        npp_pos = set(zip(pos["row"], pos["col"])); npp_neg = set(zip(neg["row"], neg["col"]))
        both = chl_pos & npp_pos; cp_nn = chl_pos & npp_neg; cp_none = chl_pos - npp_pos - npp_neg; npp_only = npp_pos - chl_pos
        row = {"year": year, "fitted": len(df), "below_floor": int((~df["floor_ok"]).sum()), "entering": len(ent),
               "retained": len(scr), "positive": len(pos), "negative": len(neg),
               "screened_window_rate": len(scr) / len(ent) if len(ent) else np.nan,
               "pos_area_km2": area(pos), "neg_area_km2": area(neg), "median_R2_fitted": float(df["R2"].median()),
               "chl_positive_cells": len(chl_pos), "positive_in_both": len(both),
               "share_positive_in_both": len(both) / len(chl_pos) if chl_pos else np.nan,
               "chl_pos_npp_neg": len(cp_nn), "chl_pos_npp_none": len(cp_none), "npp_only_positive": len(npp_only)}
        rows.append(row)
        print(f"[{year}] NPP fitted {len(df):,} | entering {len(ent):,} | retained {len(scr):,} (pos {len(pos):,} / neg {len(neg):,}) | "
              f"pos area {area(pos):,.0f} km2 | chl positive {len(chl_pos):,} | both {len(both):,} ({100 * row['share_positive_in_both']:.1f}%) | "
              f"NPP only {len(npp_only):,} | {time.time() - t0:.0f} s", flush=True)
        rows_g, cols_g = shape
        out_base = os.path.join(OUT_DIR, f"Fig_S{FIG_NUMBERS[year]}_npp_{year}_r")
        draw_r_map(scr, f"Annual Moving-window NPP–{NOX} Correlation (Pearson r), {year}; R² ≥ {R2_SCREEN}",
                   rows_g, cols_g, n_layers, lanes, out_base, have_ne)
        print(f"    [Saved] Fig. S{FIG_NUMBERS[year]}: {out_base}.{FORMATS[0]}")

        if year == REF_YEAR:
            sets = {"cpos": chl_pos, "joint": both, "disagree": cp_nn, "silent": cp_none, "npp_only": npp_only}
            aoi_name = lambda r_, c_: AOI_LIST[aoi_grid[r_, c_]][0] if aoi_grid[r_, c_] >= 0 else ""
            chl["AOI"] = [aoi_name(r_, c_) for r_, c_ in zip(chl["row"], chl["col"])]
            scr17 = scr.copy(); scr17["AOI"] = [aoi_name(r_, c_) for r_, c_ in zip(scr17["row"], scr17["col"])]
            chl.to_csv(os.path.join(WIN_DIR, f"windows_screened_chl_{year}.csv"), index=False)
            scr17.to_csv(os.path.join(WIN_DIR, f"windows_screened_npp_{year}.csv"), index=False)
            cat = [{"row": r_, "col": c_, "Class": nm, "ClassValue": v, "AOI": aoi_name(r_, c_)}
                   for nm, v, cells in (("joint", CLS_JOINT, both), ("silent", CLS_SILENT, cp_none), ("disagree", CLS_DISAGREE, cp_nn), ("npp_only", CLS_NPPONLY, npp_only))
                   for r_, c_ in sorted(cells)]
            pd.DataFrame(cat).to_csv(os.path.join(WIN_DIR, f"windows_categorized_{year}.csv"), index=False)
            t3 = table_s3(sets, aoi_grid, scr17)
            p3 = os.path.join(OUT_DIR, "Table_S3A_chl_npp_overlap_by_region.csv"); t3.to_csv(p3, index=False, encoding="utf-8-sig")
            print(f"[Saved] {p3}"); print(t3.to_string(index=False))
            aoi_gdfs = {}
            for short, _, gj in AOI_LIST:
                g = gpd.read_file(gj)
                if g.crs is not None and g.crs.to_epsg() != 4326:
                    g = g.to_crs(4326)
                aoi_gdfs[short] = g
            draw_overlap_map({CLS_JOINT: both, CLS_SILENT: cp_none, CLS_DISAGREE: cp_nn, CLS_NPPONLY: npp_only},
                             os.path.join(OUT_DIR, "Fig_S21_chl_npp_overlap"), shape, lanes, aoi_gdfs)
            checks += [f"{datetime.datetime.now():%Y-%m-%d %H:%M}", f"npp entering the screen {year}: {len(ent):,}",
                       f"classes {year}: joint {len(both):,}, silent {len(cp_none):,}, disagree {len(cp_nn):,}, npp_only {len(npp_only):,}",
                       f"npp positive on a chl-negative cell (subset of npp_only): {len(npp_pos & set(zip(chl.loc[chl['Slope'] < 0, 'row'], chl.loc[chl['Slope'] < 0, 'col']))):,}",
                       "areas of the classes, 10^6 km^2: " + "; ".join(f"{nm} {pixel_area_m2(np.array([90 - (r_ + 0.5) * yres for r_, _ in cells]), xres, yres).sum() / 1e12:.2f}"
                                                                      for nm, cells in (("joint", both), ("silent", cp_none), ("disagree", cp_nn), ("npp_only", npp_only)))]

    pd.DataFrame(rows).to_csv(os.path.join(OUT_DIR, "Table_S4_npp_interannual.csv"), index=False, encoding="utf-8-sig")
    print(f"[Saved] {os.path.join(OUT_DIR, 'Table_S4_npp_interannual.csv')}")
    tab = write_table_s4(rows, os.path.join(OUT_DIR, "Table_S4_final.csv"), os.path.join(OUT_DIR, "Table_S4_final.docx"))
    pd.set_option("display.width", 260); pd.set_option("display.max_columns", 30)
    print("\n================ TABLE S4 (paste-ready) ================"); print(tab.to_string(index=False))

    checks += table_s3b()
    npp5, emi5 = table_s5(masks)
    fig_s27(npp5, emi5, masks)
    with open(os.path.join(OUT_DIR, "checks.txt"), "w", encoding="utf-8") as fh:
        fh.write("\n".join(checks) + "\n")
    print(f"[Done] Figs. S21 to S27, Tables S3 (A, B), S4 and S5 in {OUT_DIR}/")


if __name__ == "__main__":
    main()
